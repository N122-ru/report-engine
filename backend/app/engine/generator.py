"""Generate DOCX documents according to GOST 7.32-2017 / GOST 2.105-2019."""

import io
from typing import Dict, Optional
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from ..models.report import ReportMetadata, GenerateOptions
from .title_page import add_title_page
from .markdown_parser import render_markdown_to_docx
from .bibliography import format_reference
from .oxml_utils import add_page_number_to_footer

def build_docx_report(
    meta: ReportMetadata,
    markdown_content: str,
    images_dict: Optional[Dict[str, bytes]] = None,
    options: Optional[GenerateOptions] = None
) -> io.BytesIO:
    """Generate a complete DOCX file according to GOST 7.32."""
    if options is None:
        options = GenerateOptions()
    if images_dict is None:
        images_dict = {}

    doc = Document()

    # -------------------------------------------------------------
    # Russian proofing and hyphenation avoid false spelling marks and
    # excessive spacing in justified paragraphs.
    # -------------------------------------------------------------
    rPrDefault = doc.styles.element.xpath('.//w:docDefaults/w:rPrDefault/w:rPr')
    if rPrDefault:
        for old in rPrDefault[0].findall(qn('w:lang')):
            rPrDefault[0].remove(old)
        rPrDefault[0].append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ru-RU" w:eastAsia="ru-RU" w:bidi="ar-SA"/>'))

        for old in rPrDefault[0].findall(qn('w:rFonts')):
            rPrDefault[0].remove(old)
        rPrDefault[0].append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'))

        for old in rPrDefault[0].findall(qn('w:sz')):
            rPrDefault[0].remove(old)
        rPrDefault[0].append(parse_xml(f'<w:sz {nsdecls("w")} w:val="28"/>'))

        for old in rPrDefault[0].findall(qn('w:szCs')):
            rPrDefault[0].remove(old)
        rPrDefault[0].append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="28"/>'))

    # Enable automatic hyphenation for GOST 7.32 / GOST 2.105 layout.
    doc.settings.element.append(parse_xml(f'<w:autoHyphenation {nsdecls("w")}/>'))
    doc.settings.element.append(parse_xml(f'<w:consecutiveHyphenLimit {nsdecls("w")} w:val="2"/>'))
    doc.settings.element.append(parse_xml(f'<w:hyphenationZone {nsdecls("w")} w:val="720"/>'))

    compat_el = doc.settings.element.find(qn('w:compat'))
    if compat_el is not None:
        compat_el.append(parse_xml(f'<w:doNotExpandShiftReturn {nsdecls("w")}/>'))
    # Do not set proofState to clean: it would hide the author's spelling errors.

    # -------------------------------------------------------------
    # Configure the base Normal style.
    # -------------------------------------------------------------
    normal_style = doc.styles["Normal"]
    normal_font = normal_style.font
    normal_font.name = "Times New Roman"
    normal_font.size = Pt(14)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(0)
    normal_style.paragraph_format.first_line_indent = Mm(12.5)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    rPr_normal = normal_style.element.get_or_add_rPr()
    for old in rPr_normal.findall(qn('w:lang')):
        rPr_normal.remove(old)
    rPr_normal.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ru-RU" w:eastAsia="ru-RU" w:bidi="ar-SA"/>'))

    # -------------------------------------------------------------
    # A4 margins under GOST 7.32: left 30, right 15, top/bottom 20 mm.
    # -------------------------------------------------------------
    sec_title = doc.sections[0]
    sec_title.page_width = Mm(210)
    sec_title.page_height = Mm(297)
    sec_title.left_margin = Mm(30)
    sec_title.right_margin = Mm(15)
    sec_title.top_margin = Mm(20)
    sec_title.bottom_margin = Mm(20)
    sec_title.different_first_page_header_footer = True

    # Explicit options take precedence over metadata formatting settings.
    def _opt(opt_val, meta_val):
        return opt_val if opt_val is not None else meta_val

    page_numbering = _opt(options.page_numbering, getattr(meta, "page_numbering", "center"))
    heading_alignment = _opt(options.heading_alignment, getattr(meta, "heading_alignment", "center"))
    page_break_h1 = _opt(options.page_break_before_h1, getattr(meta, "page_break_before_h1", False))
    hide_captions = _opt(options.hide_captions, getattr(meta, "hide_captions", False))

    # -------------------------------------------------------------
    # Title page.
    # -------------------------------------------------------------
    if options.include_title_page:
        add_title_page(doc, meta)

        # Create a body section only when the report contains body text.
        if markdown_content and markdown_content.strip():
            sec_body = doc.add_section(WD_SECTION_START.NEW_PAGE)
            sec_body.page_width = Mm(210)
            sec_body.page_height = Mm(297)
            sec_body.left_margin = Mm(30)
            sec_body.right_margin = Mm(15)
            sec_body.top_margin = Mm(20)
            sec_body.bottom_margin = Mm(20)

            # Unlink title-page footers and enable numbering from page two.
            sec_body.different_first_page_header_footer = False
            sec_body.footer.is_linked_to_previous = False
            if page_numbering != "none":
                add_page_number_to_footer(sec_body.footer, alignment=page_numbering)
    else:
        # Body-only output.
        sec_title.different_first_page_header_footer = False
        if page_numbering != "none":
            add_page_number_to_footer(sec_title.footer, alignment=page_numbering)

    # -------------------------------------------------------------
    # Render Markdown with one GOST-compliant listing and table style.
    # -------------------------------------------------------------
    latex_mode = _opt(options.latex_mode, getattr(meta, "latex_mode", "omml"))
    if markdown_content and markdown_content.strip():
        render_markdown_to_docx(
            doc=doc,
            markdown_text=markdown_content,
            images_dict=images_dict,
            page_break_before_h1=page_break_h1,
            heading_alignment=heading_alignment,
            hide_captions=hide_captions,
            latex_mode=latex_mode,
            image_layouts={name: layout.model_dump() for name, layout in meta.image_layouts.items()}
        )

    # -------------------------------------------------------------
    # Append GOST R 7.0.5 references only when body text exists, so a
    # title-page-only export remains one page.
    # -------------------------------------------------------------
    references = getattr(meta, "references", None) or []
    if references and markdown_content and markdown_content.strip():
        from docx.enum.style import WD_STYLE_TYPE
        try:
            bib_style = doc.styles["Bibliography"]
        except KeyError:
            bib_style = doc.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
            bib_style.base_style = doc.styles["Normal"]
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER if heading_alignment == "center" else WD_ALIGN_PARAGRAPH.LEFT
        p_head_pr = p_head._p.get_or_add_pPr()
        p_head_pr.append(parse_xml(f'<w:pageBreakBefore {nsdecls("w")}/>'))
        p_head_pr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
        p_head_pr.append(parse_xml(f'<w:keepLines {nsdecls("w")}/>'))
        p_head_pr.append(parse_xml(f'<w:suppressAutoHyphens {nsdecls("w")}/>'))
        p_head_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="0" w:left="0" w:right="0"/>'))
        p_head_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="120" w:line="360" w:lineRule="auto"/>'))
        r_head = p_head.add_run("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
        r_head.font.name = "Times New Roman"
        r_head.font.size = Pt(14)
        r_head.bold = True

        for idx, ref in enumerate(references, start=1):
            ref_dict = ref.model_dump() if hasattr(ref, "model_dump") else dict(ref)
            p_ref = doc.add_paragraph(style=bib_style)
            p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_ref_pr = p_ref._p.get_or_add_pPr()
            p_ref_pr.append(parse_xml(f'<w:keepLines {nsdecls("w")}/>'))
            p_ref_pr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
            p_ref_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="709" w:left="0" w:right="0"/>'))
            p_ref_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="0" w:line="360" w:lineRule="auto"/>'))
            r_ref = p_ref.add_run(f"{idx}.\u00A0{format_reference(ref_dict)}")
            r_ref.font.name = "Times New Roman"
            r_ref.font.size = Pt(14)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
