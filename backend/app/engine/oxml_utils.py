"""Provide low-level OpenXML helpers for precise Word formatting."""

from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt

def set_cell_background(cell, fill_hex: str):
    """Set a table cell's background color, for example F4F4F4."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex.replace("#", "")}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top: int = 100, bottom: int = 100, left: int = 150, right: int = 150):
    """Set table cell padding in dxa."""
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_table_row_cant_split(row):
    """Prevent a table row from splitting across pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_table_header_row(row):
    """Repeat a table header row on each new page."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def set_cell_bottom_border(cell, color: str = "000000", sz: str = "4"):
    """Set a cell's bottom border for a signature line."""
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(tcBorders)

def set_cell_no_borders(cell):
    """Remove all cell borders."""
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(tcBorders)

def set_table_gost_borders(table):
    """Set the standard thin GOST table borders."""
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'</w:tblBorders>'
    )
    table._tbl.tblPr.append(tblBorders)

def set_listing_clean_borders(table, has_border: bool = True):
    """Set a listing's thin single black border or remove its borders."""
    b_val = "single" if has_border else "none"
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{b_val}" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="{b_val}" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="{b_val}" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="{b_val}" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    table._tbl.tblPr.append(tblBorders)

def add_page_number_to_footer(footer, alignment: str = "center"):
    """Add a centered or right-aligned footer page number in Times New Roman 12 pt."""
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.alignment = 2 if alignment == "right" else 1 # 2: Right, 1: Center
    
    # Reset paragraph indents for footer alignment.
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="0" w:left="0" w:right="0"/>'))
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'))
    
    run = p.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    
    # Format the PAGE field explicitly as Times New Roman 12 pt.
    fld1 = parse_xml(
        f'<w:fldSimple {nsdecls("w")} w:instr="PAGE">'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
        f'      <w:sz w:val="24"/>'
        f'      <w:szCs w:val="24"/>'
        f'    </w:rPr>'
        f'    <w:t>1</w:t>'
        f'  </w:r>'
        f'</w:fldSimple>'
    )
    p._p.append(fld1)
