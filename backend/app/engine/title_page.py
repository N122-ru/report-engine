"""Generate title pages under GOST 7.32 / GOST 2.105.

Use paragraph properties and an anchored logo, not empty paragraphs, for layout.
Read institutional details from metadata, with placeholders for empty fields.
"""

import base64
import os
import re
import tempfile
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from lxml import etree

from ..models.report import ReportMetadata, StudentInfo, TeacherInfo
from .typography import strip_trailing_dot
from .oxml_utils import (
    set_cell_no_borders,
    set_cell_bottom_border,
    set_table_row_cant_split
)

candidate_paths = [
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "assets", "default_logo.png")),
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "assets", "default_logo.png")),
    os.path.abspath("assets/default_logo.png"),
]
LOGO_PATH = next((p for p in candidate_paths if os.path.exists(p)), candidate_paths[0])


def normalize_logo_mode(value: str | None) -> str:
    normalized = (value or "none").strip().lower()
    return normalized if normalized in {"default", "custom", "none"} else "none"


def _resolve_custom_logo(custom_logo: str | None) -> str | None:
    """Resolve a custom logo from a base64 data URL or file path."""
    if not custom_logo or not custom_logo.strip():
        return None
    raw = custom_logo.strip()
    # Data-URL: data:image/png;base64,....
    if "," in raw and ";base64," in raw:
        try:
            _, b64 = raw.split(";base64,", 1)
            data = base64.b64decode(b64)
        except Exception:
            return None
        suffix = ".png"
        low_head = raw[:40].lower()
        if "jpeg" in low_head or "jpg" in low_head:
            suffix = ".jpg"
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
        try:
            tmp.write(data)
            tmp.close()
            return tmp.name
        except Exception:
            return None
    if os.path.exists(raw):
        return raw
    return None


def normalize_fio_initials_first(fio: str) -> str:
    """Normalize names to initials followed by surname under GOST 7.32 / R 7.0.97.

    Accept surname-first initials, full names, and spaced initials.
    """
    fio = fio.strip()
    if not fio:
        return fio

    def _cap_name(name_str: str) -> str:
        return "-".join(part.capitalize() for part in name_str.split("-"))

    # Move surname-first initials ahead of the surname.
    m_last_first = re.match(r"^([А-ЯЁа-яё\-]+)\s+([А-ЯЁ]\.?)\s*([А-ЯЁ]\.?)?$", fio)
    if m_last_first:
        last = _cap_name(m_last_first.group(1))
        i1 = m_last_first.group(2).rstrip('.').upper() + '.'
        i2 = (m_last_first.group(3).rstrip('.').upper() + '.') if m_last_first.group(3) else ''
        return f"{i1}{i2} {last}"

    # Join initials without internal spaces when they precede the surname.
    m_first_last = re.match(r"^([А-ЯЁ]\.?)\s*([А-ЯЁ]\.?)?\s+([А-ЯЁа-яё\-]+)$", fio)
    if m_first_last:
        i1 = m_first_last.group(1).rstrip('.').upper() + '.'
        i2 = (m_first_last.group(2).rstrip('.').upper() + '.') if m_first_last.group(2) else ''
        last = _cap_name(m_first_last.group(3))
        return f"{i1}{i2} {last}"

    # Reduce full given names and patronymics to initials.
    parts = fio.split()
    if len(parts) == 3 and not any('.' in p for p in parts):
        last, first, middle = parts[0], parts[1], parts[2]
        return f"{first[0].upper()}.{middle[0].upper()}. {_cap_name(last)}"

    return fio


def add_title_page(doc: Document, meta: ReportMetadata):
    """Add the title page as the document's first page."""
    logo_mode = normalize_logo_mode(getattr(meta, "logo_mode", "none"))
    logo_path = LOGO_PATH if logo_mode == "default" else None
    if logo_mode == "custom":
        logo_path = _resolve_custom_logo(getattr(meta, "custom_logo", None))
    has_logo = bool(logo_path) and os.path.exists(logo_path)
    shift_dxa = "1276" if has_logo else "0"

    ministry_text = getattr(meta, "ministry", "") or "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ \nРОССИЙСКОЙ ФЕДЕРАЦИИ"
    if "\n" not in ministry_text:
        if "РОССИЙСКОЙ ФЕДЕРАЦИИ" in ministry_text:
            ministry_text = ministry_text.replace("РОССИЙСКОЙ ФЕДЕРАЦИИ", "\nРОССИЙСКОЙ ФЕДЕРАЦИИ").strip()
        elif len(ministry_text) > 45:
            # Split the ministry header into two balanced lines.
            parts = ministry_text.split()
            mid = len(parts) // 2
            ministry_text = " ".join(parts[:mid]) + "\n" + " ".join(parts[mid:])

    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0_pr = p0._p.get_or_add_pPr()
    p0_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="{shift_dxa}" w:firstLine="0"/>'))
    p0_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'))

    r0 = p0.add_run(ministry_text)
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(12)
    r0.bold = False

    # Anchor the enabled logo to the left of the header.
    if has_logo:
        assert logo_path is not None
        r_img = p0.add_run()
        r_img.add_picture(logo_path, width=Mm(51.4), height=Mm(27.6))
        inline = r_img._r.xpath(".//wp:inline")[0]
        extent = inline.xpath(".//wp:extent")[0]
        docPr = inline.xpath(".//wp:docPr")[0]
        graphic = inline.xpath(".//a:graphic")[0]

        anchor_xml = f'''<wp:anchor {nsdecls("wp")} {nsdecls("a")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" hidden="0" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="column">
            <wp:posOffset>-932179</wp:posOffset>
          </wp:positionH>
          <wp:positionV relativeFrom="paragraph">
            <wp:posOffset>-292099</wp:posOffset>
          </wp:positionV>
          {etree.tostring(extent).decode()}
          <wp:effectExtent l="0" t="0" r="0" b="0"/>
          <wp:wrapNone/>
          {etree.tostring(docPr).decode()}
          <wp:cNvGraphicFramePr/>
          {etree.tostring(graphic).decode()}
        </wp:anchor>'''
        drawing = inline.getparent()
        drawing.remove(inline)
        drawing.append(parse_xml(anchor_xml))

    # Use metadata or a placeholder for the university name.
    # Reserve header space for the floating logo when enabled.
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1_pr = p1._p.get_or_add_pPr()
    p1_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="{shift_dxa}" w:firstLine="0"/>'))
    if shift_dxa != "0":
        p1_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'))
    else:
        p1_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="0" w:line="240" w:lineRule="auto"/>'))
    uni_title = meta.university.strip() if meta.university and meta.university.strip() else "НАИМЕНОВАНИЕ УНИВЕРСИТЕТА"
    r1 = p1.add_run(uni_title)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.bold = False

    # Optional institute or faculty.
    if meta.institute and meta.institute.strip():
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3_pr = p3._p.get_or_add_pPr()
        p3_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="{shift_dxa}" w:firstLine="0"/>'))
        p3_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="120" w:after="0" w:line="240" w:lineRule="auto"/>'))
        r3 = p3.add_run(meta.institute)
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(12)
        r3.bold = False

    # Optional department.
    if meta.department and meta.department.strip():
        p4 = doc.add_paragraph()
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p4_pr = p4._p.get_or_add_pPr()
        p4_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="{shift_dxa}" w:firstLine="0"/>'))
        p4_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'))
        clean_dept = re.sub(r'^(кафедра|кафедры)\s*', '', meta.department.strip(), flags=re.IGNORECASE).strip('"«»')
        r4 = p4.add_run(f"Кафедра «{clean_dept}»")
        r4.font.name = "Times New Roman"
        r4.font.size = Pt(12)
        r4.bold = False

    # -------------------------------------------------------------
    # Estimate title-page spacing to target a single-page layout.
    # -------------------------------------------------------------
    num_teachers = len(meta.teachers) if meta.teachers is not None else 1
    num_students = len(meta.students) if meta.students is not None else 1
    num_persons = num_teachers + num_students
    total_rows = num_persons * 2

    theme_lines = max(1, (len(meta.work_theme) + 35) // 45)
    compact_mode = num_persons > 3 or theme_lines > 2
    
    # Estimate element heights in dxa: ministry/university header,
    # followed by optional institute and department paragraphs.
    if has_logo:
        header_dxa = 1710 if (meta.institute or meta.department) else 1110
    else:
        header_dxa = 1500 if (meta.institute or meta.department) else 1050
    title_fixed = 1180 + (theme_lines * 300)
    person_row_height = 750
    table_dxa = num_persons * person_row_height
    footer_dxa = 240

    fixed_dxa = header_dxa + title_fixed + table_dxa + footer_dxa
    target_total = 14250 - max(0, num_persons - 2) * 150 - max(0, theme_lines - 1) * 120
    avail_dxa = max(0, target_total - fixed_dxa)

    if avail_dxa > 1000:
        gap1 = int(avail_dxa * 0.30)
        gap2 = int(avail_dxa * 0.32)
        gap3 = max(180, avail_dxa - gap1 - gap2)
    elif avail_dxa > 0:
        gap1 = max(100, int(avail_dxa * 0.30))
        gap2 = max(100, int(avail_dxa * 0.32))
        gap3 = max(140, avail_dxa - gap1 - gap2)
    else:
        gap1 = 40
        gap2 = 40
        gap3 = 60

    # -------------------------------------------------------------
    # Report heading block.
    # -------------------------------------------------------------
    p_rep = doc.add_paragraph()
    p_rep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rep_pr = p_rep._p.get_or_add_pPr()
    p_rep_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="{gap1}" w:after="120" w:line="240" w:lineRule="auto"/>'))
    p_rep_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="0" w:left="0" w:right="0"/>'))
    r_rep = p_rep.add_run(meta.doc_type)
    r_rep.font.name = "Times New Roman"
    r_rep.font.size = Pt(18)
    r_rep.bold = True

    p_work = doc.add_paragraph()
    p_work.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_work_pr = p_work._p.get_or_add_pPr()
    p_work_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="80" w:line="240" w:lineRule="auto"/>'))
    p_work_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="0" w:left="0" w:right="0"/>'))
    work_num_str = f" №{meta.work_number}" if (meta.work_number is not None and meta.work_number > 0) else ""
    r_work = p_work.add_run(f"{meta.work_type}{work_num_str}")
    r_work.font.name = "Times New Roman"
    r_work.font.size = Pt(13)

    clean_theme = strip_trailing_dot(meta.work_theme)
    if not clean_theme.startswith("«"):
        clean_theme = f"«{clean_theme}»"

    p_theme = doc.add_paragraph()
    p_theme.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_theme_pr = p_theme._p.get_or_add_pPr()
    p_theme_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="80" w:line="240" w:lineRule="auto"/>'))
    p_theme_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="0" w:left="0" w:right="0"/>'))
    r_theme = p_theme.add_run(clean_theme)
    r_theme.font.name = "Times New Roman"
    r_theme.font.size = Pt(14)
    r_theme.bold = True

    p_disc = doc.add_paragraph()
    p_disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_disc_pr = p_disc._p.get_or_add_pPr()
    p_disc_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="{gap2}" w:line="240" w:lineRule="auto"/>'))
    p_disc_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="0" w:left="0" w:right="0"/>'))
    clean_disc = meta.discipline.strip().strip('"').strip('«').strip('»')
    r_disc = p_disc.add_run(f"по дисциплине «{clean_disc}»")
    r_disc.font.name = "Times New Roman"
    r_disc.font.size = Pt(13)

    # -------------------------------------------------------------
    # Signature table.
    # -------------------------------------------------------------
    sig_table = doc.add_table(rows=total_rows, cols=5)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    # Column widths sum to the 165 mm text area.
    col_widths = [Mm(37), Mm(30), Mm(30), Mm(25), Mm(43)]
    for row in sig_table.rows:
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]

    # Keep signature-table padding compact.
    tbl_pr = sig_table._tbl.tblPr
    tbl_cell_mar = parse_xml(f'''
        <w:tblCellMar {nsdecls("w")}>
            <w:top w:w="30" w:type="dxa"/>
            <w:left w:w="40" w:type="dxa"/>
            <w:bottom w:w="30" w:type="dxa"/>
            <w:right w:w="40" w:type="dxa"/>
        </w:tblCellMar>
    ''')
    tbl_pr.append(tbl_cell_mar)

    sub_size = Pt(8)
    cell_spacing = 10 if compact_mode else 20

    current_row = 0
    # An explicit empty list means no signatories; do not insert phantom names.
    # Model defaults apply only when the field is omitted.
    teachers = meta.teachers if meta.teachers is not None else [TeacherInfo(fio="", role="")]
    teacher_label = "Преподаватели" if len(teachers) > 1 else "Преподаватель"
    for t_idx, teacher in enumerate(teachers):
        r_data = sig_table.rows[current_row]
        r_sub = sig_table.rows[current_row + 1]
        set_table_row_cant_split(r_data)
        set_table_row_cant_split(r_sub)

        _set_cell_text(r_data.cells[0], teacher_label if t_idx == 0 else "", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, size=Pt(12), spacing=cell_spacing)
        _set_cell_text(r_data.cells[1], teacher.role, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11), spacing=cell_spacing)
        _set_cell_text(r_data.cells[2], "", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11), spacing=cell_spacing)
        _set_cell_text(r_data.cells[3], "", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11), spacing=cell_spacing)
        _set_cell_text(r_data.cells[4], normalize_fio_initials_first(teacher.fio), bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11), spacing=cell_spacing)

        set_cell_no_borders(r_data.cells[0])
        set_cell_bottom_border(r_data.cells[1])
        set_cell_bottom_border(r_data.cells[2])
        set_cell_bottom_border(r_data.cells[3])
        set_cell_bottom_border(r_data.cells[4])

        _set_cell_text(r_sub.cells[0], "", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)
        _set_cell_text(r_sub.cells[1], "(должность)", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)
        _set_cell_text(r_sub.cells[2], "(подпись)", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)
        _set_cell_text(r_sub.cells[3], "(дата)", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)
        _set_cell_text(r_sub.cells[4], "(инициалы, фамилия)", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)

        for c in r_sub.cells:
            set_cell_no_borders(c)

        current_row += 2

    students = meta.students if meta.students is not None else [StudentInfo(fio="", group="")]
    student_label = "Студенты" if len(students) > 1 else "Студент"
    for s_idx, student in enumerate(students):
        r_data = sig_table.rows[current_row]
        r_sub = sig_table.rows[current_row + 1]
        set_table_row_cant_split(r_data)
        set_table_row_cant_split(r_sub)

        _set_cell_text(r_data.cells[0], student_label if s_idx == 0 else "", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, size=Pt(12), spacing=cell_spacing)
        _set_cell_text(r_data.cells[1], student.group, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11), spacing=cell_spacing)
        _set_cell_text(r_data.cells[2], "", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11), spacing=cell_spacing)
        _set_cell_text(r_data.cells[3], "", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11), spacing=cell_spacing)
        _set_cell_text(r_data.cells[4], normalize_fio_initials_first(student.fio), bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11), spacing=cell_spacing)

        set_cell_no_borders(r_data.cells[0])
        set_cell_bottom_border(r_data.cells[1])
        set_cell_bottom_border(r_data.cells[2])
        set_cell_bottom_border(r_data.cells[3])
        set_cell_bottom_border(r_data.cells[4])

        _set_cell_text(r_sub.cells[0], "", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)
        _set_cell_text(r_sub.cells[1], "(группа)", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)
        _set_cell_text(r_sub.cells[2], "(подпись)", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)
        _set_cell_text(r_sub.cells[3], "(дата)", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)
        _set_cell_text(r_sub.cells[4], "(инициалы, фамилия)", size=sub_size, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=0)

        for c in r_sub.cells:
            set_cell_no_borders(c)

        current_row += 2

    # -------------------------------------------------------------
    # Footer: city and year without a year abbreviation.
    # -------------------------------------------------------------
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer_pr = p_footer._p.get_or_add_pPr()
    p_footer_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="{gap3}" w:after="0" w:line="240" w:lineRule="auto"/>'))
    p_footer_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="0" w:left="0" w:right="0"/>'))
    clean_city = re.sub(r"^\s*г\.\s*|^\s*г\s+", "", (meta.city or "").strip(), flags=re.IGNORECASE)
    clean_city = re.sub(r"[\s,]+г\.?\s*$", "", clean_city, flags=re.IGNORECASE).strip() or "Город"
    r_footer = p_footer.add_run(f"{clean_city} {meta.year}")
    r_footer.font.name = "Times New Roman"
    r_footer.font.size = Pt(12)


def _set_cell_text(cell, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, size=Pt(11), spacing: int = 20):
    """Set table cell text using the required formatting."""
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.text = ""
    p.alignment = align
    p_pr = p._p.get_or_add_pPr()
    p_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="{spacing}" w:after="{spacing}" w:line="240" w:lineRule="auto"/>'))
    p_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="0" w:left="0" w:right="0"/>'))

    if text:
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = size
        run.bold = bold
