import sys, os
sys.path.insert(0, os.path.abspath("backend"))

import json
import pytest
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def _list_all_mcp_tools(token):
    tools = []
    cursor = None
    request_id = 200
    while True:
        params = {} if cursor is None else {"cursor": cursor}
        response = client.post(
            f"/api/mcp/messages?token={token}",
            json={"jsonrpc": "2.0", "id": request_id, "method": "tools/list", "params": params},
        )
        assert response.status_code == 200
        result = response.json()["result"]
        tools.extend(result["tools"])
        cursor = result.get("nextCursor")
        request_id += 1
        if not cursor:
            return tools


def test_session_init_creates_new_token():
    res = client.get("/api/session/init")
    assert res.status_code == 200
    data = res.json()
    assert "token" in data
    assert data["token"].startswith("tok_")
    assert data["content"] == ""


def test_session_sync_and_get():
    # Create a session.
    res1 = client.get("/api/session/init")
    token = res1.json()["token"]

    # Synchronize browser state.
    sync_res = client.post(
        f"/api/session/{token}/sync",
        json={"content": "# Ход работы\nТекст пользователя", "meta": {"work_number": 5}}
    )
    assert sync_res.status_code == 200

    # Verify session restoration.
    res2 = client.get(f"/api/session/init?token={token}")
    data2 = res2.json()
    assert data2["token"] == token
    assert "Ход работы" in data2["content"]
    assert data2["meta"]["work_number"] == 5


def test_session_sync_accepts_compact_unicode_content_patch():
    token = client.get("/api/session/init").json()["token"]
    original = "Параметр: 🙂\nСтарое значение"
    initial = client.post(f"/api/session/{token}/sync", json={"content": original}).json()
    patch = client.post(f"/api/session/{token}/sync", json={
        "content_patches": [{"start": 12, "delete_count": 15, "insert": "Новое значение"}],
        "expected_revision": initial["revision"],
    })
    assert patch.status_code == 200
    assert client.get(f"/api/session/init?token={token}").json()["content"] == "Параметр: 🙂\nНовое значение"


@pytest.mark.anyio
async def test_large_content_event_uses_delta_instead_of_full_snapshot():
    import asyncio
    from app.engine.sessions import SessionManager

    manager = SessionManager()
    session = manager.get_or_create()
    events = asyncio.Queue()
    session.add_listener(events)
    original = "x" * 1200
    await manager.update_content(session.token, original, source="mcp_agent")
    await events.get()
    await manager.update_content(session.token, original + " tail", source="mcp_agent")
    event = await events.get()
    session.remove_listener(events)
    assert "content_patches" in event
    assert "content" not in event
    assert event["content_base_revision"]


def test_browser_sync_rejects_stale_content_without_losing_agent_edit():
    token = client.get("/api/session/init").json()["token"]
    initial = client.get(f"/api/session/init?token={token}").json()
    assert initial["revision"]
    first = client.post(f"/api/session/{token}/sync", json={
        "content": "# Ход работы\nПравка браузера", "expected_revision": initial["revision"]
    })
    assert first.status_code == 200
    first_revision = first.json()["revision"]
    _, live = _call_report_tool(token, "get_report_state", {})
    result, _ = _call_report_tool(token, "append_or_update_section", {
        "section_heading": "# Вывод", "section_content": "Правка агента",
        "mode": "update_or_append", "expected_revision": live["revision"]
    })
    assert result["isError"] is False

    stale = client.post(f"/api/session/{token}/sync", json={
        "content": "# Ход работы\nЗапоздалый черновик", "expected_revision": first_revision
    })
    assert stale.status_code == 409
    current = client.get(f"/api/session/init?token={token}").json()
    assert stale.json()["revision"] == current["revision"]
    assert stale.json()["content"] == current["content"]
    assert "Правка агента" in current["content"]


def test_browser_sync_rejects_stale_metadata_without_losing_agent_edit():
    token = client.get("/api/session/init").json()["token"]
    initial = client.get(f"/api/session/init?token={token}").json()
    assert initial["meta_revision"]
    first = client.post(f"/api/session/{token}/sync", json={
        "meta": {"work_theme": "Черновик"},
        "expected_meta_revision": initial["meta_revision"]
    })
    assert first.status_code == 200
    result, _ = _call_report_tool(token, "update_meta", {
        "work_theme": "Тема от агента", "expected_meta_revision": first.json()["meta_revision"]
    })
    assert result["isError"] is False
    stale = client.post(f"/api/session/{token}/sync", json={
        "meta": {"work_theme": "Старое название"},
        "expected_meta_revision": first.json()["meta_revision"]
    })
    assert stale.status_code == 409
    assert stale.json()["conflict_type"] == "meta"
    assert stale.json()["meta"]["work_theme"] == "Тема от агента"
    assert client.get(f"/api/session/init?token={token}").json()["meta"]["work_theme"] == "Тема от агента"


@pytest.mark.anyio
async def test_browser_sync_marks_sse_event_with_originating_client():
    from app.engine.sessions import session_manager
    import asyncio

    token = client.get("/api/session/init").json()["token"]
    session = session_manager.get_or_create(token)
    events = asyncio.Queue()
    session.add_listener(events)
    await session_manager.update_content(token, "# Ход работы\nТекст.", source="web", client_id="tab-one")
    event = await events.get()
    session.remove_listener(events)
    assert event["type"] == "content_updated"
    assert event["source"] == "web"
    assert event["client_id"] == "tab-one"


@pytest.mark.anyio
async def test_unchanged_metadata_does_not_broadcast_another_event():
    import asyncio
    from app.engine.sessions import SessionManager

    manager = SessionManager()
    session = manager.get_or_create()
    events = asyncio.Queue()
    session.add_listener(events)
    meta = {"work_theme": "Тема", "image_layouts": {"screen.png": {"align": "center", "width_percent": 72}}}
    await manager.update_meta(session.token, meta, source="mcp_agent")
    revision = session.meta_revision
    assert (await events.get())["type"] == "meta_updated"

    await manager.update_meta(session.token, {
        "image_layouts": {"screen.png": {"width_percent": 72, "align": "center"}}, "work_theme": "Тема",
    }, source="mcp_agent")
    assert session.meta_revision == revision
    assert events.empty()


def test_repeated_agent_metadata_patch_reports_unchanged():
    token = client.get("/api/session/init").json()["token"]
    _, state = _call_report_tool(token, "get_report_state", {})
    _, first = _call_report_tool(token, "update_meta", {
        "work_theme": "Тема", "expected_meta_revision": state["meta_revision"],
    })
    _, state = _call_report_tool(token, "get_report_state", {})
    result, repeated = _call_report_tool(token, "update_meta", {
        "work_theme": "Тема", "expected_meta_revision": state["meta_revision"],
    })
    assert result["isError"] is False
    assert repeated["status"] == "unchanged"
    assert repeated["meta_revision"] == state["meta_revision"] == first["meta_revision"]


def test_agent_meta_patch_with_stale_revision_is_rejected():
    token = client.get("/api/session/init").json()["token"]
    _, state = _call_report_tool(token, "get_report_state", {})
    client.post(f"/api/session/{token}/sync", json={
        "meta": {"work_theme": "Свежий текст пользователя"},
        "expected_meta_revision": state["meta_revision"]
    })
    result, _ = _call_report_tool(token, "update_meta", {
        "work_theme": "Устаревшая версия", "expected_meta_revision": state["meta_revision"]
    })
    assert result["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["meta"]["work_theme"] == "Свежий текст пользователя"


def test_mcp_mutations_require_revisions_even_in_default_modes():
    token = client.get("/api/session/init").json()["token"]
    original = "# Вывод\nТекст студента.\n"
    client.post(f"/api/session/{token}/sync", json={"content": original, "meta": {"work_theme": "Тема студента"}})

    section, _ = _call_report_tool(token, "append_or_update_section", {
        "section_heading": "# Вывод", "section_content": "Текст агента."
    })
    meta, _ = _call_report_tool(token, "update_meta", {"work_theme": "Тема агента"})
    assert section["isError"] is True
    assert meta["isError"] is True
    state = client.get(f"/api/session/init?token={token}").json()
    assert state["content"] == original
    assert state["meta"]["work_theme"] == "Тема студента"

    tools = _list_all_mcp_tools(token)
    schemas = {tool["name"]: tool["inputSchema"] for tool in tools}
    assert "expected_revision" in schemas["append_or_update_section"]["required"]
    assert "expected_meta_revision" in schemas["update_meta"]["required"]
    assert "get_images" in schemas
    assert "delete_images" in schemas
    assert "expected_image_revision" in schemas["delete_image"]["required"]


def test_agent_cannot_rollback_last_browser_edit():
    token = client.get("/api/session/init").json()["token"]
    initial = client.get(f"/api/session/init?token={token}").json()
    edited = client.post(f"/api/session/{token}/sync", json={
        "content": "# Ход работы\nНовый текст студента.", "expected_revision": initial["revision"]
    }).json()
    result, _ = _call_report_tool(token, "rollback_last_edit", {
        "expected_revision": edited["revision"]
    })
    assert result["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["content"] == "# Ход работы\nНовый текст студента."


def test_mcp_autofix_preserves_markdown_links_and_code_and_requires_revision():
    token = client.get("/api/session/init").json()["token"]
    original = (
        "# Ход работы.\n\nТекст - описание.\n\n"
        "Снимок ![Схема](Wi-Fi.png) и [файл](https://example.org/Wi-Fi.png).\n\n"
        "~~~bash\nWi-Fi.png и IP-адрес\n~~~\n\n"
        "    Wi-Fi.png и IP-адрес\n\n"
        "Команда `Wi-Fi.png` не меняется.\n"
    )
    client.post(f"/api/session/{token}/sync", json={"content": original})
    missing, _ = _call_report_tool(token, "autofix_typography", {})
    assert missing["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["content"] == original

    _, state = _call_report_tool(token, "get_report_state", {})
    result, payload = _call_report_tool(token, "autofix_typography", {
        "expected_revision": state["revision"]
    })
    assert result["isError"] is False
    current = client.get(f"/api/session/init?token={token}").json()["content"]
    assert current.startswith("# Ход работы\n\nТекст\u00a0— описание.")
    assert "Снимок ![Схема](Wi-Fi.png) и [файл](https://example.org/Wi-Fi.png)." in current
    assert "~~~bash\nWi-Fi.png и IP-адрес\n~~~" in current
    assert "    Wi-Fi.png и IP-адрес" in current
    assert "`Wi-Fi.png`" in current
    assert payload["revision"] != state["revision"]
    assert payload["diff"]

    stale, _ = _call_report_tool(token, "autofix_typography", {
        "expected_revision": state["revision"]
    })
    assert stale["isError"] is True


def test_remote_mcp_handshake_and_tools():
    # Create a session.
    res1 = client.get("/api/session/init")
    token = res1.json()["token"]

    # MCP initialize
    init_res = client.post(
        f"/api/mcp/messages?token={token}",
        json={"jsonrpc": "2.0", "id": 1, "method": "initialize"}
    )
    assert init_res.status_code == 200
    init_data = init_res.json()
    assert init_data["result"]["serverInfo"]["name"] == "report-engine-live"
    assert init_data["result"]["serverInfo"]["sessionToken"] == token

    # MCP tools/list, including all pagination pages.
    tool_names = [tool["name"] for tool in _list_all_mcp_tools(token)]
    assert "get_report_state" in tool_names
    assert "append_or_update_section" in tool_names
    assert "lint_report" in tool_names
    assert "update_meta" in tool_names
    assert "delete_image" in tool_names
    assert "upload_image" in tool_names
    assert "import_profile" in tool_names
    assert "set_full_report" not in tool_names
    assert "get_templates" not in tool_names


def test_mcp_tools_list_keeps_core_on_first_page_and_paginates_the_rest():
    token = client.get("/api/session/init").json()["token"]

    first = client.post(
        f"/api/mcp/messages?token={token}",
        json={"jsonrpc": "2.0", "id": 2, "method": "tools/list"},
    )
    assert first.status_code == 200
    first_result = first.json()["result"]
    assert len(first_result["tools"]) == 10
    assert first_result["nextCursor"]
    first_names = {tool["name"] for tool in first_result["tools"]}
    assert {
        "get_report_state", "list_images", "get_image", "get_images", "upload_image",
        "append_or_update_section", "lint_report", "autofix_typography", "update_meta", "delete_image",
    } <= first_names

    all_tools = _list_all_mcp_tools(token)
    all_names = [tool["name"] for tool in all_tools]
    assert len(all_names) == len(set(all_names))
    assert set(all_names) == {
        "list_images", "get_image", "get_images", "upload_image", "upload_images",
        "update_image_layout", "get_report_state", "append_or_update_section", "patch_text",
        "apply_patches", "get_document_outline", "replace_block", "delete_block", "move_figure",
        "get_last_edit", "get_edit_history", "rollback_last_edit", "lint_report",
        "autofix_typography", "update_meta", "delete_image", "delete_images", "import_profile",
        "batch",
    }


@pytest.mark.anyio
async def test_parallel_mcp_writes_are_serialized_by_revision():
    import asyncio
    from app.engine.mcp_router import _dispatch_mcp_tool_safely
    from app.engine.sessions import session_manager

    session = session_manager.get_or_create()
    expected_revision = session.revision

    results = await asyncio.gather(
        _dispatch_mcp_tool_safely(
            "append_or_update_section",
            {
                "section_heading": "# Цель работы",
                "section_content": "Первая параллельная запись.",
                "expected_revision": expected_revision,
            },
            session,
        ),
        _dispatch_mcp_tool_safely(
            "append_or_update_section",
            {
                "section_heading": "# Вывод",
                "section_content": "Вторая параллельная запись.",
                "expected_revision": expected_revision,
            },
            session,
        ),
        return_exceptions=True,
    )

    assert sum(not isinstance(result, Exception) for result in results) == 1
    assert sum(isinstance(result, ValueError) for result in results) == 1
    assert "параллельная запись" in session.content


def test_mcp_sse_post_response_is_not_returned_twice():
    import asyncio
    from app.engine.mcp_router import _mcp_connections

    token = client.get("/api/session/init").json()["token"]
    connection_id = "test-no-duplicate"
    queue = asyncio.Queue()
    _mcp_connections[connection_id] = queue
    try:
        response = client.post(
            f"/api/mcp/messages?token={token}&sessionId={connection_id}",
            json={"jsonrpc": "2.0", "id": 22, "method": "ping"},
        )
        assert response.status_code == 202
        queued = queue.get_nowait()
        assert queued["id"] == 22
    finally:
        _mcp_connections.pop(connection_id, None)


def test_mcp_update_meta_merges_and_validates():
    res1 = client.get("/api/session/init")
    token = res1.json()["token"]
    client.post(
        f"/api/session/{token}/sync",
        json={"content": "# Цель работы\nТекст.", "meta": {"university": "Старый вуз", "city": "Москва"}}
    )
    current = client.get(f"/api/session/init?token={token}").json()

    # A targeted update changes only the supplied fields.
    upd = client.post(
        f"/api/mcp/messages?token={token}",
        json={
            "jsonrpc": "2.0",
            "id": 11,
            "method": "tools/call",
            "params": {
                "name": "update_meta",
                "arguments": {"work_theme": "Новая тема", "teachers": [{"fio": "И.И. Иванов", "role": "доцент"}],
                              "expected_meta_revision": current["meta_revision"]}
            }
        }
    )
    assert upd.status_code == 200
    assert upd.json()["result"]["isError"] is False
    meta = client.get(f"/api/session/init?token={token}").json()["meta"]
    assert meta["work_theme"] == "Новая тема"
    assert meta["university"] == "Старый вуз"
    assert meta["city"] == "Москва"

    # Reject empty patches.
    bad = client.post(
        f"/api/mcp/messages?token={token}",
        json={
            "jsonrpc": "2.0",
            "id": 12,
            "method": "tools/call",
            "params": {"name": "update_meta", "arguments": {}}
        }
    )
    assert bad.json()["result"]["isError"] is True


def test_mcp_import_profile_applies_both_subject_and_title_requisites_atomically():
    token = client.get("/api/session/init").json()["token"]
    state = client.get(f"/api/session/init?token={token}").json()
    profile = {
        "format": "report-engine-profile",
        "version": 1,
        "subject": {
            "id": "subj_profile",
            "name": "Сетевые технологии",
            "institute": "Институт технологий",
            "department": "Кафедра сетей",
            "teachers": [{"fio": "И.И. Иванов", "role": "доцент"}],
            "custom_instructions": "Сохранять команды.",
            "last_work_number": 4,
        },
        "meta": {
            "university": "Учебный университет",
            "ministry": "Министерство образования",
            "logo_mode": "none",
            "city": "Москва",
            "institute": "Институт технологий",
            "department": "Кафедра сетей",
            "doc_type": "ОТЧЕТ",
            "work_type": "о выполнении лабораторной работы",
            "work_number": 5,
            "discipline": "Сетевые технологии",
            "students": [{"fio": "П.П. Петров", "group": "ИВТ-101"}],
            "teachers": [{"fio": "И.И. Иванов", "role": "доцент"}],
            "year": 2026,
        },
    }
    result, payload = _call_report_tool(token, "import_profile", {
        "profile": profile,
        "expected_meta_revision": state["meta_revision"],
    })

    assert result["isError"] is False
    assert payload["subject"]["name"] == "Сетевые технологии"
    imported = client.get(f"/api/session/init?token={token}").json()
    assert imported["meta"]["discipline"] == "Сетевые технологии"
    assert imported["meta"]["custom_instructions"] == "Сохранять команды."
    assert imported["meta_revision"] != state["meta_revision"]

    stale, _ = _call_report_tool(token, "import_profile", {
        "profile": profile,
        "expected_meta_revision": state["meta_revision"],
    })
    assert stale["isError"] is True


def test_mcp_delete_image_tool():
    res1 = client.get("/api/session/init")
    token = res1.json()["token"]
    sample_png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    client.post(
        f"/api/session/{token}/images",
        files={"files": ("doomed.png", sample_png, "image/png")}
    )
    assert len(client.get(f"/api/session/init?token={token}").json()["images"]) == 1
    current = client.get(f"/api/session/init?token={token}").json()

    rm = client.post(
        f"/api/mcp/messages?token={token}",
        json={
            "jsonrpc": "2.0",
            "id": 13,
            "method": "tools/call",
            "params": {"name": "delete_image", "arguments": {
                "filename": "doomed.png", "expected_revision": current["revision"],
                "expected_image_revision": current["image_revision"]
            }}
        }
    )
    assert rm.json()["result"]["isError"] is False
    import json as _json
    rm_payload = _json.loads(rm.json()["result"]["content"][0]["text"])
    assert rm_payload["status"] == "deleted"
    assert len(client.get(f"/api/session/init?token={token}").json()["images"]) == 0


def test_mcp_batch_delete_images_is_atomic_and_skips_referenced_files():
    token = client.get("/api/session/init").json()["token"]
    sample = b"image bytes"
    upload = client.post(
        f"/api/session/{token}/images",
        files=[
            ("files", ("keep.png", sample, "image/png")),
            ("files", ("remove.png", sample, "image/png")),
        ],
    )
    assert upload.status_code == 200
    client.post(f"/api/session/{token}/sync", json={"content": "![Keep](keep.png)"})
    state = client.get(f"/api/session/init?token={token}").json()

    blocked, _ = _call_report_tool(token, "delete_images", {
        "filenames": ["keep.png", "remove.png"],
        "expected_revision": state["revision"],
        "expected_image_revision": state["image_revision"],
    })
    assert blocked["isError"] is True
    assert {item["filename"] for item in client.get(f"/api/session/{token}/images").json()["images"]} == {
        "keep.png", "remove.png"
    }

    deleted, payload = _call_report_tool(token, "delete_images", {
        "filenames": ["remove.png"],
        "expected_revision": state["revision"],
        "expected_image_revision": state["image_revision"],
    })
    assert deleted["isError"] is False
    assert payload["deleted"] == ["remove.png"]
    assert [item["filename"] for item in client.get(f"/api/session/{token}/images").json()["images"]] == ["keep.png"]


def test_stale_mcp_image_revision_cannot_delete_replacement():
    token = client.get("/api/session/init").json()["token"]
    client.post(f"/api/session/{token}/images", files={"files": ("screen.png", b"first", "image/png")})
    first = client.get(f"/api/session/init?token={token}").json()
    removed = client.delete(
        f"/api/session/{token}/images/screen.png?expected_image_revision={first['image_revision']}"
    )
    assert removed.status_code == 200
    client.post(f"/api/session/{token}/images", files={"files": ("screen.png", b"replacement", "image/png")})
    result, _ = _call_report_tool(token, "delete_image", {
        "filename": "screen.png",
        "expected_revision": first["revision"],
        "expected_image_revision": first["image_revision"],
    })
    assert result["isError"] is True
    assert client.get(f"/api/session/{token}/images/screen.png").content == b"replacement"


def test_session_accepts_150_images():
    from app.engine.sessions import MAX_IMAGES_COUNT, SessionData

    session = SessionData("tok_media_limit")
    for index in range(MAX_IMAGES_COUNT):
        session.add_image(f"image_{index}.png", b"file")

    assert len(session.images) == 150


def test_skill_guide_resource_serves_skill_md():
    import os
    res1 = client.get("/api/session/init")
    token = res1.json()["token"]
    res = client.post(
        f"/api/mcp/messages?token={token}",
        json={
            "jsonrpc": "2.0",
            "id": 14,
            "method": "resources/read",
            "params": {"uri": "report://rules/gost732"}
        }
    )
    assert res.status_code == 200
    text = res.json()["result"]["contents"][0]["text"]
    assert "безличная форма" in text.lower()
    assert os.path.exists("assets/SKILL.md")


def test_remote_mcp_updates_live_session():
    # Create a session with initial text.
    res1 = client.get("/api/session/init")
    token = res1.json()["token"]
    client.post(
        f"/api/session/{token}/sync",
        json={"content": "# Цель работы\nЦель исследования.", "meta": {}}
    )
    current = client.get(f"/api/session/init?token={token}").json()

    # The agent calls append_or_update_section.
    mcp_call = client.post(
        f"/api/mcp/messages?token={token}",
        json={
            "jsonrpc": "2.0",
            "id": 10,
            "method": "tools/call",
            "params": {
                "name": "append_or_update_section",
                "arguments": {
                    "section_heading": "# Ход работы",
                    "section_content": "1. Выполнена настройка стенда.",
                    "expected_revision": current["revision"]
                }
            }
        }
    )
    assert mcp_call.status_code == 200
    call_res = mcp_call.json()
    assert "result" in call_res
    assert call_res["result"]["isError"] is False

    # Verify that session text updates immediately.
    check_res = client.get(f"/api/session/init?token={token}")
    session_data = check_res.json()
    assert "# Ход работы" in session_data["content"]
    assert "1. Выполнена настройка стенда." in session_data["content"]
    assert "# Цель работы" in session_data["content"]


def _call_report_tool(token, name, arguments):
    response = client.post(
        f"/api/mcp/messages?token={token}",
        json={"jsonrpc": "2.0", "id": 77, "method": "tools/call",
              "params": {"name": name, "arguments": arguments}},
    )
    assert response.status_code == 200
    result = response.json()["result"]
    return result, json.loads(result["content"][0]["text"]) if not result["isError"] else None


def test_mcp_batch_runs_mutation_and_status_calls_in_order():
    token = client.get("/api/session/init").json()["token"]
    _, state = _call_report_tool(token, "get_report_state", {})
    result, payload = _call_report_tool(token, "batch", {"calls": [
        {"name": "update_meta", "arguments": {
            "work_theme": "Проверка", "expected_meta_revision": state["meta_revision"],
        }},
        {"name": "get_report_state", "arguments": {}},
        {"name": "list_images", "arguments": {}},
    ]})
    assert result["isError"] is False
    assert all(item["ok"] for item in payload)
    assert [item["name"] for item in payload] == ["update_meta", "get_report_state", "list_images"]
    assert payload[1]["result"]["metadata"]["work_theme"] == "Проверка"
    assert payload[2]["result"]["images"] == []


@pytest.mark.parametrize("stop_on_error, count", [(False, 2), (True, 1)])
def test_mcp_batch_collects_error_and_optionally_stops(stop_on_error, count):
    token = client.get("/api/session/init").json()["token"]
    result, payload = _call_report_tool(token, "batch", {"stop_on_error": stop_on_error, "calls": [
        {"name": "unknown", "arguments": {}}, {"name": "get_report_state", "arguments": {}},
    ]})
    assert result["isError"] is False
    assert len(payload) == count
    assert payload[0]["ok"] is False
    assert "Неизвестный инструмент" in payload[0]["error"]
    if not stop_on_error:
        assert payload[1]["ok"] is True


@pytest.mark.parametrize("calls", [
    [], [{"name": "list_images", "arguments": {}}] * 21,
    [{"name": "list_images", "arguments": []}],
    [{"name": "update_meta", "arguments": {"work_theme": "Не записывать"}},
     {"name": "batch", "arguments": {}}],
])
def test_mcp_invalid_batch_rejects_entire_array(calls):
    token = client.get("/api/session/init").json()["token"]
    before = client.get(f"/api/session/init?token={token}").json()
    result, _ = _call_report_tool(token, "batch", {"calls": calls})
    assert result["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json() == before


@pytest.mark.anyio
async def test_mcp_batch_holds_write_lock_for_entire_sequence_and_allows_reads():
    import asyncio
    from app.engine.mcp_router import _dispatch_mcp_tool_safely
    from app.engine.sessions import session_manager

    session = session_manager.get_or_create()
    events = asyncio.Queue(maxsize=1)
    session.add_listener(events)
    await events.put({"type": "occupied"})
    batch = asyncio.create_task(_dispatch_mcp_tool_safely("batch", {"calls": [
        {"name": "update_meta", "arguments": {
            "work_theme": "Первая тема", "expected_meta_revision": session.meta_revision,
        }},
        {"name": "get_report_state", "arguments": {}},
    ]}, session))
    competing = None
    try:
        await asyncio.sleep(0)
        assert session.mcp_write_lock.locked()
        read = await _dispatch_mcp_tool_safely("batch", {"calls": [
            {"name": "get_report_state", "arguments": {}},
        ]}, session)
        assert read[0]["result"]["metadata"]["work_theme"] == "Первая тема"
        competing = asyncio.create_task(_dispatch_mcp_tool_safely("update_meta", {
            "work_theme": "Вторая тема", "expected_meta_revision": session.meta_revision,
        }, session))
        await asyncio.sleep(0)
        assert not competing.done()
        await events.get()
        results = await asyncio.wait_for(batch, timeout=2)
        assert results[1]["result"]["metadata"]["work_theme"] == "Первая тема"
        await events.get()
        await asyncio.wait_for(competing, timeout=2)
        assert session.meta["work_theme"] == "Вторая тема"
    finally:
        session.remove_listener(events)
        for task in (batch, competing):
            if task is not None and not task.done():
                task.cancel()
                await asyncio.gather(task, return_exceptions=True)


def test_mcp_upload_image_reads_file_path_and_inserts_markdown(tmp_path):
    token = client.get("/api/session/init").json()["token"]
    _, state = _call_report_tool(token, "get_report_state", {})
    sample = b"\x89PNG\r\n\x1a\n" + b"payload"
    path = tmp_path / "screen.png"
    path.write_bytes(sample)
    result, payload = _call_report_tool(token, "upload_image", {
        "filename": "Снимок экрана.png", "file_path": str(path),
        "expected_revision": state["revision"], "caption": "Снимок", "insert_into_report": True,
    })
    assert result["isError"] is False
    assert payload["filename"] == "Снимок_экрана.png"
    assert client.get(f"/api/session/{token}/images/Снимок_экрана.png").content == sample
    assert "![Снимок](Снимок_экрана.png)" in client.get(f"/api/session/init?token={token}").json()["content"]


@pytest.mark.parametrize("source, message", [
    ({}, "base64_data или file_path"),
    ({"file_path": "screen.png", "base64_data": "AAAA"}, "base64_data или file_path"),
    ({"file_path": "missing.png"}, "не найден"),
])
def test_mcp_upload_image_rejects_invalid_file_sources(tmp_path, source, message):
    token = client.get("/api/session/init").json()["token"]
    _, state = _call_report_tool(token, "get_report_state", {})
    if "file_path" in source:
        source = {**source, "file_path": str(tmp_path / source["file_path"])}
    result, _ = _call_report_tool(token, "upload_image", {
        "filename": "screen.png", "expected_revision": state["revision"], **source,
    })
    assert result["isError"] is True
    assert message in result["content"][0]["text"]
    assert client.get(f"/api/session/{token}/images").json()["images"] == []


def test_mcp_upload_images_accepts_files_and_remains_atomic(tmp_path):
    from app.engine.sessions import MAX_IMAGE_SIZE

    token = client.get("/api/session/init").json()["token"]
    _, state = _call_report_tool(token, "get_report_state", {})
    path = tmp_path / "one.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"payload")
    large = tmp_path / "large.png"
    with large.open("wb") as output:
        output.truncate(MAX_IMAGE_SIZE + 1)
    invalid, _ = _call_report_tool(token, "upload_images", {
        "expected_revision": state["revision"], "images": [
            {"filename": "one.png", "file_path": str(path)},
            {"filename": "large.png", "file_path": str(large)},
        ],
    })
    assert invalid["isError"] is True
    assert "15 МБ" in invalid["content"][0]["text"]
    assert client.get(f"/api/session/{token}/images").json()["images"] == []
    valid, payload = _call_report_tool(token, "upload_images", {
        "expected_revision": state["revision"], "images": [
            {"filename": "one.png", "file_path": str(path)},
            {"filename": "two.png", "file_path": str(path)},
        ],
    })
    assert valid["isError"] is False
    assert payload["uploaded_count"] == 2
    schemas = {tool["name"]: tool["inputSchema"] for tool in _list_all_mcp_tools(token)}
    assert schemas["batch"]["properties"]["calls"]["maxItems"] == 20
    assert "base64_data" not in schemas["upload_image"]["required"]
    assert "file_path" in schemas["upload_images"]["properties"]["images"]["items"]["properties"]


@pytest.mark.anyio
@pytest.mark.parametrize("tool", ["upload_image", "upload_images"])
async def test_mcp_file_reads_leave_loop_responsive_and_recheck_revision(tmp_path, monkeypatch, tool):
    import asyncio
    import threading
    from app.engine import mcp_router
    from app.engine.sessions import session_manager

    session = session_manager.get_or_create()
    path = tmp_path / "screen.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\n")
    started = threading.Event()
    release = threading.Event()
    original = mcp_router.load_image_input

    def delayed_read(args, decoder):
        started.set()
        if not release.wait(timeout=3):
            raise ValueError("Чтение файла заблокировало цикл событий")
        return original(args, decoder)

    monkeypatch.setattr(mcp_router, "load_image_input", delayed_read)
    item = {"filename": "screen.png", "file_path": str(path)}
    args = {"expected_revision": session.revision, **(item if tool == "upload_image" else {"images": [item]})}
    upload = asyncio.create_task(mcp_router._dispatch_mcp_tool_safely(tool, args, session))
    try:
        assert await asyncio.to_thread(started.wait, 2)
        await session_manager.update_content(session.token, "Правка браузера.", source="web")
        release.set()
        with pytest.raises(ValueError, match="Документ изменился"):
            await asyncio.wait_for(upload, timeout=2)
        assert session.images == {}
        assert session.content == "Правка браузера."
    finally:
        release.set()
        if not upload.done():
            upload.cancel()
        await asyncio.gather(upload, return_exceptions=True)


def test_section_replace_preserves_nested_headings_code_and_siblings():
    token = client.get("/api/session/init").json()["token"]
    original = (
        "# Цель работы\nЦель.\n\n# Ход работы\nСтарый текст.\n\n"
        "## Задание 1\nСтарое задание.\n\n"
        "```text\n# Это строка кода\n```\n\n"
        "# Вывод\nВывод сохраняется.\n"
    )
    client.post(f"/api/session/{token}/sync", json={"content": original})
    state, state_data = _call_report_tool(token, "get_report_state", {})
    assert state["isError"] is False
    revision = state_data["revision"]
    args = {
        "section_heading": "# Ход работы",
        "section_content": "Обновлено.\n\n## Задание 1\nНовый текст.",
        "mode": "replace_body",
        "expected_revision": revision,
    }
    result, payload = _call_report_tool(token, "append_or_update_section", args)
    assert result["isError"] is False
    updated = client.get(f"/api/session/init?token={token}").json()["content"]
    assert updated == (
        "# Цель работы\nЦель.\n\n# Ход работы\n\nОбновлено.\n\n"
        "## Задание 1\nНовый текст.\n\n# Вывод\nВывод сохраняется.\n"
    )
    assert payload["headings_before"] == ["# Цель работы", "# Ход работы", "## Задание 1", "# Вывод"]
    assert payload["headings_after"] == ["# Цель работы", "# Ход работы", "## Задание 1", "# Вывод"]
    assert payload["revision"] != revision
    assert "Вывод сохраняется" in payload["diff"] or "Обновлено." in payload["diff"]


def test_section_replace_ignores_headings_in_blockquotes():
    token = client.get("/api/session/init").json()["token"]
    original = "# Ход работы\nНачало.\n\n> # Вывод\n> Это лишь цитата.\n\n# Вывод\nРеальный итог.\n"
    client.post(f"/api/session/{token}/sync", json={"content": original})
    _, state = _call_report_tool(token, "get_report_state", {})
    result, _ = _call_report_tool(token, "append_or_update_section", {
        "section_heading": "# Ход работы", "section_content": "Новая работа.",
        "mode": "replace_body", "expected_revision": state["revision"],
    })
    assert result["isError"] is False
    assert client.get(f"/api/session/init?token={token}").json()["content"] == (
        "# Ход работы\n\nНовая работа.\n\n# Вывод\nРеальный итог.\n"
    )


def test_section_append_to_body_and_duplicate_heading_rejection():
    token = client.get("/api/session/init").json()["token"]
    original = "# Ход работы\nПервое.\n\n## Задание 1\nТекст.\n\n# Вывод\nИтог.\n"
    client.post(f"/api/session/{token}/sync", json={"content": original})
    _, state = _call_report_tool(token, "get_report_state", {})
    result, payload = _call_report_tool(token, "append_or_update_section", {
        "section_heading": "# Ход работы", "section_content": "## Задание 2\nДобавлено.",
        "mode": "append_to_body", "expected_revision": state["revision"]
    })
    assert result["isError"] is False
    assert payload["status"] == "updated"
    expected = "# Ход работы\nПервое.\n\n## Задание 1\nТекст.\n\n## Задание 2\nДобавлено.\n\n# Вывод\nИтог.\n"
    assert client.get(f"/api/session/init?token={token}").json()["content"] == expected
    stale_result, _ = _call_report_tool(token, "append_or_update_section", {
        "section_heading": "# Вывод", "section_content": "Чужие изменения.",
        "mode": "replace_body", "expected_revision": state["revision"]
    })
    assert stale_result["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["content"] == expected

    client.post(f"/api/session/{token}/sync", json={"content": "# Вывод\nПервый.\n\n# Вывод\nВторой.\n"})
    _, duplicate_state = _call_report_tool(token, "get_report_state", {})
    duplicate_result, _ = _call_report_tool(token, "append_or_update_section", {
        "section_heading": "# Вывод", "section_content": "Замена.",
        "mode": "replace_body", "expected_revision": duplicate_state["revision"]
    })
    assert duplicate_result["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["content"] == "# Вывод\nПервый.\n\n# Вывод\nВторой.\n"


def test_new_section_rejects_repeated_heading_inside_content():
    token = client.get("/api/session/init").json()["token"]
    original = "# Цель работы\nПлан.\n"
    client.post(f"/api/session/{token}/sync", json={"content": original})
    _, state = _call_report_tool(token, "get_report_state", {})
    result, _ = _call_report_tool(token, "append_or_update_section", {
        "section_heading": "# Вывод", "section_content": "# Вывод\nИтог.",
        "mode": "update_or_append", "expected_revision": state["revision"],
    })
    assert result["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["content"] == original


def test_section_id_edits_only_selected_duplicate_heading():
    token = client.get("/api/session/init").json()["token"]
    original = ("# Ход работы\n\n## Задание\nПервое.\n\n"
                "## Задание\nВторое.\n\n# Вывод\nСохранить.\n")
    client.post(f"/api/session/{token}/sync", json={"content": original})
    _, outline = _call_report_tool(token, "get_document_outline", {})
    targets = [section for section in outline["sections"] if section["heading"] == "## Задание"]
    assert len(targets) == 2
    result, _ = _call_report_tool(token, "append_or_update_section", {
        "section_heading": "## Задание", "section_id": targets[1]["id"],
        "section_content": "Новое второе.", "mode": "replace_body",
        "expected_revision": outline["revision"],
    })
    assert result["isError"] is False
    current = client.get(f"/api/session/init?token={token}").json()["content"]
    assert "## Задание\nПервое." in current
    assert "## Задание\n\nНовое второе." in current
    assert "Второе." not in current
    assert "# Вывод\nСохранить." in current


def test_mcp_linter_distinguishes_review_from_invalid():
    token = client.get("/api/session/init").json()["token"]
    client.post(f"/api/session/{token}/sync", json={
        "content": "# Цель работы\nВыполнено исследование.\n\n# Ход работы\nСнимок представлен ниже.\n\n![Рисунок 4.2 — Схема](missing.png)\n\n# Вывод\nПолучен результат."
    })
    result, payload = _call_report_tool(token, "lint_report", {})
    assert result["isError"] is False
    assert payload["is_valid"] is True  # Legacy behavior has no blocking errors.
    assert payload["status"] == "needs_review"
    assert any(w["type"] == "IMAGES" for w in payload["warnings"])


def test_list_images_reports_dimensions_and_actual_markdown_insertion():
    from app.engine.sessions import SessionData
    session = SessionData("tok_dimension_test")
    png = (b"\x89PNG\r\n\x1a\n" + b"\x00" * 8 +
           (320).to_bytes(4, "big") + (180).to_bytes(4, "big") + b"\x00" * 12)
    session.add_image("placed.png", png)
    session.add_image("mentioned.png", png)
    session.content = ("Упомянут (mentioned.png).\n\n"
                       "```markdown\n![Скрин](mentioned.png)\n```\n\n"
                       "![Рабочий скрин](placed.png)")
    images = {image["filename"]: image for image in session.list_images()}
    assert images["placed.png"]["is_in_report"] is True
    assert images["mentioned.png"]["is_in_report"] is False
    assert images["placed.png"]["width_px"] == 320
    assert images["placed.png"]["height_px"] == 180


def test_cyrillic_image_name_stays_addressable_and_is_detected_in_markdown():
    from app.engine.sessions import SessionData
    session = SessionData("tok_cyrillic_image")
    name = session.add_image("Схема сети.png", b"payload")
    session.content = "![Рисунок 1 — Схема](<Схема сети.png>)"
    assert name == "Схема_сети.png"
    assert session.get_image("Схема сети.png") == b"payload"
    assert session.list_images()[0]["is_in_report"] is True


def test_patch_text_replaces_one_unique_phrase_atomically():
    token = client.get("/api/session/init").json()["token"]
    client.post(f"/api/session/{token}/sync", json={"content": "# Ход работы\nПодпись рисунка: схема стенда.\n\n# Вывод\nГотово.\n"})
    _, state = _call_report_tool(token, "get_report_state", {})
    result, payload = _call_report_tool(token, "patch_text", {
        "find": "Подпись рисунка: схема стенда.",
        "replace": "Подпись рисунка: топология стенда.",
        "expected_revision": state["revision"],
    })
    assert result["isError"] is False
    assert payload["replacements"] == 1
    assert "Подпись рисунка: топология стенда." in client.get(f"/api/session/init?token={token}").json()["content"]

    stale, _ = _call_report_tool(token, "patch_text", {
        "find": "Готово.", "replace": "Изменено.", "expected_revision": state["revision"]
    })
    assert stale["isError"] is True


def test_batch_patch_applies_together_and_rolls_back_as_one_edit():
    token = client.get("/api/session/init").json()["token"]
    original = "# Ход работы\nКод на рисунке 1.\nРезультат на рисунке 2.\n\n# Вывод\nИтог.\n"
    client.post(f"/api/session/{token}/sync", json={"content": original})
    _, state = _call_report_tool(token, "get_report_state", {})
    invalid, _ = _call_report_tool(token, "apply_patches", {
        "expected_revision": state["revision"],
        "patches": [{"find": "Код на рисунке 1.", "replace": "Код приведён на рисунке 1."},
                    {"find": "Несуществующий абзац", "replace": "Чужой текст"}]
    })
    assert invalid["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["content"] == original
    valid, payload = _call_report_tool(token, "apply_patches", {
        "expected_revision": state["revision"],
        "patches": [{"find": "Код на рисунке 1.", "replace": "Код приведён на рисунке 1."},
                    {"find": "Результат на рисунке 2.", "replace": "Результат проверки — на рисунке 2."}]
    })
    assert valid["isError"] is False
    assert payload["replacements"] == 2
    _, history = _call_report_tool(token, "get_edit_history", {})
    assert history["edits"][0]["before_revision"] == state["revision"]
    assert history["edits"][0]["after_revision"] == payload["revision"]
    result, _ = _call_report_tool(token, "rollback_last_edit", {"expected_revision": payload["revision"]})
    assert result["isError"] is False
    assert client.get(f"/api/session/init?token={token}").json()["content"] == original


def test_last_edit_diff_and_conditional_rollback():
    token = client.get("/api/session/init").json()["token"]
    original = "# Ход работы\nИсходный текст.\n"
    client.post(f"/api/session/{token}/sync", json={"content": original})
    _, state = _call_report_tool(token, "get_report_state", {})
    changed, _ = _call_report_tool(token, "patch_text", {
        "find": "Исходный текст.", "replace": "Обновлённый текст.",
        "expected_revision": state["revision"]
    })
    assert changed["isError"] is False
    _, last = _call_report_tool(token, "get_last_edit", {})
    assert "Обновлённый текст." in last["diff"]
    assert last["source"] == "mcp_agent"
    _, updated = _call_report_tool(token, "get_report_state", {})
    stale, _ = _call_report_tool(token, "rollback_last_edit", {"expected_revision": state["revision"]})
    assert stale["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["content"] != original
    result, restored = _call_report_tool(token, "rollback_last_edit", {"expected_revision": updated["revision"]})
    assert result["isError"] is False
    assert restored["revision"] == state["revision"]
    assert client.get(f"/api/session/init?token={token}").json()["content"] == original


def test_replace_and_delete_markdown_blocks_by_snapshot_id():
    token = client.get("/api/session/init").json()["token"]
    original = ("# Ход работы\nОписание до рисунка.\n\n"
                "![Рисунок 4.2 — Схема](a.png)\n\n# Вывод\nИтог сохраняется.\n")
    client.post(f"/api/session/{token}/sync", json={"content": original})
    _, outline = _call_report_tool(token, "get_document_outline", {})
    image = next(block for block in outline["blocks"] if block["type"] == "image")
    paragraph = next(block for block in outline["blocks"] if block["type"] == "paragraph" and "Описание" in block["preview"])
    result, changed = _call_report_tool(token, "replace_block", {
        "block_id": image["id"],
        "content": "![Рисунок 4.2 — Новая схема](b.png)",
        "expected_revision": outline["revision"],
    })
    assert result["isError"] is False
    assert "b.png" in client.get(f"/api/session/init?token={token}").json()["content"]
    stale, _ = _call_report_tool(token, "delete_block", {
        "block_id": paragraph["id"], "expected_revision": outline["revision"]
    })
    assert stale["isError"] is True
    _, new_outline = _call_report_tool(token, "get_document_outline", {})
    paragraph = next(block for block in new_outline["blocks"] if block["type"] == "paragraph" and "Описание" in block["preview"])
    deleted, _ = _call_report_tool(token, "delete_block", {
        "block_id": paragraph["id"], "expected_revision": new_outline["revision"]
    })
    assert deleted["isError"] is False
    final = client.get(f"/api/session/init?token={token}").json()["content"]
    assert "Описание до рисунка" not in final
    assert "Итог сохраняется." in final
    assert "Новая схема" in final
    assert changed["revision"] == new_outline["revision"]


def test_move_figure_before_another_block_without_duplication():
    token = client.get("/api/session/init").json()["token"]
    original = "# Ход работы\n\n![Рисунок 1 — A](a.png)\n\n![Рисунок 2 — B](b.png)\n\n# Вывод\nИтог.\n"
    client.post(f"/api/session/{token}/sync", json={"content": original})
    _, outline = _call_report_tool(token, "get_document_outline", {})
    figures = [block for block in outline["blocks"] if block["type"] == "image"]
    result, _ = _call_report_tool(token, "move_figure", {
        "block_id": figures[1]["id"], "before_block_id": figures[0]["id"],
        "expected_revision": outline["revision"]
    })
    assert result["isError"] is False
    text = client.get(f"/api/session/init?token={token}").json()["content"]
    assert text.index("b.png") < text.index("a.png")
    assert text.count("b.png") == text.count("a.png") == 1
    assert "# Вывод\nИтог." in text


def test_mcp_image_layout_updates_metadata_and_rejects_stale_version():
    token = client.get("/api/session/init").json()["token"]
    from app.engine.sessions import session_manager
    session_manager.get_or_create(token).add_image("diagram.png", b"file")
    _, state = _call_report_tool(token, "get_report_state", {})
    result, layout = _call_report_tool(token, "update_image_layout", {
        "filename": "diagram.png", "width_percent": 72, "align": "right",
        "max_height_mm": 120, "page_break_before": True,
        "expected_meta_revision": state["meta_revision"]
    })
    assert result["isError"] is False
    assert layout["layout"]["width_percent"] == 72
    meta = client.get(f"/api/session/init?token={token}").json()["meta"]
    assert meta["image_layouts"]["diagram.png"]["align"] == "right"
    stale, _ = _call_report_tool(token, "update_image_layout", {
        "filename": "diagram.png", "width_percent": 80,
        "expected_meta_revision": state["meta_revision"]
    })
    assert stale["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["meta"]["image_layouts"]["diagram.png"]["width_percent"] == 72


def test_batch_upload_images_inserts_at_anchor_and_rejects_partial_invalid_batch():
    import base64
    token = client.get("/api/session/init").json()["token"]
    original = "# Ход работы\nПодготовлен стенд.\n\n# Вывод\nИтог.\n"
    client.post(f"/api/session/{token}/sync", json={"content": original})
    _, outline = _call_report_tool(token, "get_document_outline", {})
    anchor = next(b for b in outline["blocks"] if b["type"] == "paragraph" and "Подготовлен" in b["preview"])
    encoded = base64.b64encode(b"picture").decode()
    invalid, _ = _call_report_tool(token, "upload_images", {
        "expected_revision": outline["revision"],
        "images": [
            {"filename": "one.png", "base64_data": encoded, "caption": "Код", "after_block_id": anchor["id"]},
            {"filename": "two.png", "base64_data": "bad-base64", "caption": "Вывод"}
        ]
    })
    assert invalid["isError"] is True
    assert client.get(f"/api/session/init?token={token}").json()["content"] == original
    assert client.get(f"/api/session/{token}/images").json()["images"] == []
    result, payload = _call_report_tool(token, "upload_images", {
        "expected_revision": outline["revision"],
        "images": [
            {"filename": "one.png", "base64_data": encoded, "caption": "Рисунок 1 — Код", "after_block_id": anchor["id"]},
            {"filename": "two.png", "base64_data": encoded, "caption": "Рисунок 2 — Вывод", "after_block_id": anchor["id"]}
        ]
    })
    assert result["isError"] is False
    assert payload["uploaded_count"] == 2
    content = client.get(f"/api/session/init?token={token}").json()["content"]
    assert content.index("Подготовлен") < content.index("one.png") < content.index("two.png") < content.index("# Вывод")
    assert len(client.get(f"/api/session/{token}/images").json()["images"]) == 2


def test_single_mcp_upload_rejects_missing_revision_bad_base64_and_existing_name():
    import base64
    token = client.get("/api/session/init").json()["token"]
    _, state = _call_report_tool(token, "get_report_state", {})
    original = b"original screenshot"
    client.post(f"/api/session/{token}/images", files={"files": ("Wi-Fi.png", original, "image/png")})

    missing, _ = _call_report_tool(token, "upload_image", {
        "filename": "other.png", "base64_data": base64.b64encode(b"picture").decode()
    })
    bad, _ = _call_report_tool(token, "upload_image", {
        "filename": "other.png", "base64_data": "!!!", "expected_revision": state["revision"]
    })
    collision, _ = _call_report_tool(token, "upload_image", {
        "filename": "Wi-Fi.png", "base64_data": base64.b64encode(b"replacement").decode(),
        "expected_revision": state["revision"]
    })
    assert missing["isError"] and bad["isError"] and collision["isError"]
    assert client.get(f"/api/session/{token}/images/Wi-Fi.png").content == original
    assert len(client.get(f"/api/session/{token}/images").json()["images"]) == 1


def test_browser_image_batch_failure_leaves_no_partial_upload():
    from app.engine.sessions import MAX_IMAGE_SIZE
    token = client.get("/api/session/init").json()["token"]
    response = client.post(f"/api/session/{token}/images", files=[
        ("files", ("first.png", b"first", "image/png")),
        ("files", ("oversized.png", b"x" * (MAX_IMAGE_SIZE + 1), "image/png")),
    ])
    assert response.status_code == 400
    assert client.get(f"/api/session/{token}/images").json()["images"] == []


def test_browser_image_upload_cannot_replace_original_through_sanitized_name():
    token = client.get("/api/session/init").json()["token"]
    first = client.post(f"/api/session/{token}/images", files={
        "files": ("Схема сети.png", b"original", "image/png")
    })
    assert first.status_code == 200
    second = client.post(f"/api/session/{token}/images", files={
        "files": ("Схема_сети.png", b"processed", "image/png")
    })
    assert second.status_code == 400
    assert client.get(f"/api/session/{token}/images/Схема_сети.png").content == b"original"


def test_mcp_cannot_delete_image_still_referenced_by_markdown():
    token = client.get("/api/session/init").json()["token"]
    client.post(f"/api/session/{token}/images", files={"files": ("used.png", b"picture", "image/png")})
    client.post(f"/api/session/{token}/sync", json={"content": "![Схема](used.png)"})
    _, state = _call_report_tool(token, "get_report_state", {})
    result, _ = _call_report_tool(token, "delete_image", {
        "filename": "used.png", "expected_revision": state["revision"],
        "expected_image_revision": state["image_revision"]
    })
    assert result["isError"] is True
    assert client.get(f"/api/session/{token}/images/used.png").content == b"picture"


def test_live_session_tracks_image_deletions_until_the_server_restarts():
    token = client.get("/api/session/init").json()["token"]
    client.post(f"/api/session/{token}/images", files={
        "files": ("Схема сети.png", b"original", "image/png")
    })
    _, state = _call_report_tool(token, "get_report_state", {})
    result, _ = _call_report_tool(token, "delete_image", {
        "filename": "Схема_сети.png", "expected_revision": state["revision"],
        "expected_image_revision": state["image_revision"]
    })
    assert result["isError"] is False
    images = client.get(f"/api/session/{token}/images").json()
    assert images["images"] == []
    assert images["deleted"] == ["Схема_сети.png"]

    client.post(f"/api/session/{token}/images", files={
        "files": ("Схема_сети.png", b"new original", "image/png")
    })
    assert client.get(f"/api/session/{token}/images").json()["deleted"] == []


def test_restore_upload_cannot_undo_deletion_after_a_stale_image_list():
    token = client.get("/api/session/init").json()["token"]
    client.post(f"/api/session/{token}/images", files={
        "files": ("old.png", b"original", "image/png")
    })
    _, state = _call_report_tool(token, "get_report_state", {})
    result, _ = _call_report_tool(token, "delete_image", {
        "filename": "old.png", "expected_revision": state["revision"],
        "expected_image_revision": state["image_revision"]
    })
    assert result["isError"] is False
    stale_restore = client.post(f"/api/session/{token}/images?restoring=true", files={
        "files": ("old.png", b"original", "image/png")
    })
    assert stale_restore.status_code == 409
    assert client.get(f"/api/session/{token}/images").json()["images"] == []
    assert client.get(f"/api/session/{token}/images").json()["deleted"] == ["old.png"]


def test_mcp_linter_reports_attached_but_unused_images():
    token = client.get("/api/session/init").json()["token"]
    from app.engine.sessions import session_manager
    session = session_manager.get_or_create(token)
    session.add_image("used.png", b"file")
    session.add_image("unused.png", b"file")
    client.post(f"/api/session/{token}/sync", json={
        "content": "# Ход работы\n\n![Рисунок 1 — Тест](used.png)\n\n# Вывод\nИтог."
    })
    result, payload = _call_report_tool(token, "lint_report", {})
    assert result["isError"] is False
    assert any(w["type"] == "UNUSED_IMAGE" and "unused.png" in w["message"] for w in payload["warnings"])
    assert not any(w["type"] == "UNUSED_IMAGE" and w["location"] == "used.png" for w in payload["warnings"])


def test_mcp_linter_checks_declared_task_and_figure_counts():
    token = client.get("/api/session/init").json()["token"]
    client.post(f"/api/session/{token}/sync", json={
        "content": "# Ход работы\n1. Первое действие.\n2. Второе действие.\n\n![Рисунок 1 — Схема](a.png)\n"
    })
    result, payload = _call_report_tool(token, "lint_report", {
        "expected_task_count": 3, "expected_figures_count": 2
    })
    assert result["isError"] is False
    assert payload["metrics"]["tasks_count"] == 2
    assert payload["metrics"]["figures_count"] == 1
    assert len([w for w in payload["warnings"] if w["type"] == "COUNT_MISMATCH"]) == 2


def test_session_image_upload_and_crud():
    # Create a session.
    res1 = client.get("/api/session/init")
    token = res1.json()["token"]

    # Upload an image through the browser session endpoint.
    sample_png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    upload_res = client.post(
        f"/api/session/{token}/images",
        files=[("files", ("wireshark_capture.png", sample_png, "image/png"))]
    )
    assert upload_res.status_code == 200
    up_data = upload_res.json()
    assert len(up_data["uploaded"]) == 1
    assert up_data["uploaded"][0]["filename"] == "wireshark_capture.png"

    # Verify the image list.
    list_res = client.get(f"/api/session/{token}/images")
    assert list_res.status_code == 200
    images = list_res.json()["images"]
    assert len(images) == 1
    assert images[0]["filename"] == "wireshark_capture.png"
    import hashlib
    assert images[0]["sha256"] == hashlib.sha256(sample_png).hexdigest()

    # Download the image through its direct URL.
    get_res = client.get(f"/api/session/{token}/images/wireshark_capture.png")
    assert get_res.status_code == 200
    assert get_res.headers["content-type"].startswith("image/png")
    assert get_res.content == sample_png

    # Delete the image.
    del_res = client.delete(f"/api/session/{token}/images/wireshark_capture.png")
    assert del_res.status_code == 200
    assert del_res.json()["status"] == "deleted"

    # Verify that the image is no longer served.
    get_again = client.get(f"/api/session/{token}/images/wireshark_capture.png")
    assert get_again.status_code == 404


def test_mcp_multimodal_image_tools():
    import base64

    # Create a session and upload a screenshot.
    res1 = client.get("/api/session/init")
    token = res1.json()["token"]
    sample_png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4"
    client.post(
        f"/api/session/{token}/images",
        files=[("files", ("network_dump.png", sample_png, "image/png"))]
    )

    # The agent calls list_images.
    list_call = client.post(
        f"/api/mcp/messages?token={token}",
        json={
            "jsonrpc": "2.0",
            "id": 11,
            "method": "tools/call",
            "params": {"name": "list_images", "arguments": {}}
        }
    )
    assert list_call.status_code == 200
    list_json = json.loads(list_call.json()["result"]["content"][0]["text"])
    assert list_json["total_count"] == 1
    assert list_json["images"][0]["filename"] == "network_dump.png"

    # A multimodal agent calls get_image for visual inspection.
    get_call = client.post(
        f"/api/mcp/messages?token={token}",
        json={
            "jsonrpc": "2.0",
            "id": 12,
            "method": "tools/call",
            "params": {"name": "get_image", "arguments": {"filename": "network_dump.png"}}
        }
    )
    assert get_call.status_code == 200
    content_blocks = get_call.json()["result"]["content"]
    # The first content block must be a native MCP image.
    assert content_blocks[0]["type"] == "image"
    assert content_blocks[0]["mimeType"] == "image/png"
    assert base64.b64decode(content_blocks[0]["data"]) == sample_png

    # Batch retrieval returns multiple image blocks in one MCP call.
    batch_call = client.post(
        f"/api/mcp/messages?token={token}",
        json={
            "jsonrpc": "2.0",
            "id": 121,
            "method": "tools/call",
            "params": {"name": "get_images", "arguments": {"filenames": ["network_dump.png"]}}
        }
    )
    assert batch_call.status_code == 200
    batch_blocks = batch_call.json()["result"]["content"]
    assert batch_blocks[0]["type"] == "image"
    assert base64.b64decode(batch_blocks[0]["data"]) == sample_png

    # The agent uploads a new chart through upload_image.
    b64_graph = base64.b64encode(b"FAKE_PNG_GRAPH_DATA").decode("ascii")
    current = client.get(f"/api/session/init?token={token}").json()
    upload_call = client.post(
        f"/api/mcp/messages?token={token}",
        json={
            "jsonrpc": "2.0",
            "id": 13,
            "method": "tools/call",
            "params": {
                "name": "upload_image",
                "arguments": {
                    "filename": "generated_chart.png",
                        "base64_data": b64_graph,
                        "caption": "График зависимости времени отклика",
                        "insert_into_report": True,
                        "expected_revision": current["revision"]
                }
            }
        }
    )
    assert upload_call.status_code == 200
    up_res = json.loads(upload_call.json()["result"]["content"][0]["text"])
    assert up_res["status"] == "success"
    assert up_res["inserted_into_report"] is True

    # Verify the report contains a GOST-compatible image reference.
    state_call = client.get(f"/api/session/init?token={token}")
    state = state_call.json()
    assert "![График зависимости времени отклика](generated_chart.png)" in state["content"]


from starlette.requests import Request
from app.engine.mcp_router import mcp_sse_transport

@pytest.mark.anyio
async def test_mcp_sse_endpoint_respects_reverse_proxy_https():
    """Verify SSE endpoint events preserve the HTTPS reverse proxy's scheme and origin."""
    scope = {
        "type": "http",
        "method": "GET",
        "path": "/api/mcp/sse",
        "headers": [
            (b"x-forwarded-proto", b"https"),
            (b"x-forwarded-host", b"lab.example.com"),
            (b"host", b"lab.example.com"),
        ],
        "query_string": b"token=tok_testproxy"
    }
    req = Request(scope)
    res = await mcp_sse_transport(req, token="tok_testproxy")
    assert res.status_code == 200
    first_chunk = await res.body_iterator.__anext__()
    assert first_chunk.startswith("event: endpoint\n")
    assert "https://lab.example.com/api/mcp/messages?token=tok_testproxy" in first_chunk

    # Fall back to HTTP during local development without a proxy.
    scope_http = {
        "type": "http",
        "method": "GET",
        "path": "/api/mcp/sse",
        "headers": [
            (b"host", b"localhost:8000"),
        ],
        "query_string": b"token=tok_testproxy"
    }
    req_http = Request(scope_http)
    res_http = await mcp_sse_transport(req_http, token="tok_testproxy")
    first_chunk_http = await res_http.body_iterator.__anext__()
    assert "http://localhost:8000/api/mcp/messages?token=tok_testproxy" in first_chunk_http

    # Exercise comma-separated proxy headers, uppercase HTTPS, and a forwarded prefix.
    scope_multi = {
        "type": "http",
        "method": "GET",
        "path": "/api/mcp/sse",
        "headers": [
            (b"x-forwarded-proto", b"HTTPS, http"),
            (b"x-forwarded-host", b"lab.example.com, internal-proxy:8080"),
            (b"x-forwarded-prefix", b"/sub-report/"),
            (b"host", b"internal-proxy:8080"),
        ],
        "query_string": b"token=tok_nestedproxy"
    }
    req_multi = Request(scope_multi)
    res_multi = await mcp_sse_transport(req_multi, token="tok_nestedproxy")
    first_chunk_multi = await res_multi.body_iterator.__anext__()
    assert "https://lab.example.com/sub-report/api/mcp/messages?token=tok_nestedproxy" in first_chunk_multi


def test_token_neutrality_and_validation():
    """Verify tok_/rep_tok_ formats and invalid-token sanitization."""
    from app.engine.sessions import session_manager

    # New tokens use the default prefix.
    new_tok = session_manager.create_token()
    assert new_tok.startswith("tok_")
    assert session_manager._is_valid_token(new_tok) is True

    # Accept supported token formats.
    assert session_manager._is_valid_token("tok_abcdef1234567890") is True
    assert session_manager._is_valid_token("rep_tok_998877665544") is True

    # Reject invalid formats and injection attempts.
    assert session_manager._is_valid_token("custom_tok_123") is False
    assert session_manager._is_valid_token("poly_tok_legacy_v1") is False
    assert session_manager._is_valid_token("tok_") is False
    assert session_manager._is_valid_token("../../../etc/passwd") is False
    assert session_manager._is_valid_token("tok_evil<script>") is False

    # Replace invalid get_or_create input with a safe generated token.
    safe_session = session_manager.get_or_create("malicious_attempt_123")
    assert safe_session.token.startswith("tok_")
    assert safe_session.token != "malicious_attempt_123"


def test_report_metadata_university_neutrality():
    """Verify metadata defaults do not impose an institution or work theme."""
    from app.models.report import ReportMetadata

    meta = ReportMetadata()
    assert meta.university == ""
    assert meta.institute == ""
    assert meta.department == ""
    assert meta.logo_mode == "none"
    assert meta.work_theme == ""
    assert meta.discipline == ""


def test_title_page_header_geometry_and_modes():
    """Verify header geometry differs between none and default logo modes."""
    import docx
    from app.models.report import ReportMetadata
    from app.engine.generator import build_docx_report

    # Use a fictional university without a logo.
    meta_neutral = ReportMetadata(
        university="Государственный технический университет",
        logo_mode="none"
    )
    doc_neutral = docx.Document(build_docx_report(meta_neutral, "# Ход работы"))
    
    # The first paragraph must not contain an anchored image.
    anchors_neutral = doc_neutral.paragraphs[0]._p.xpath('.//*[local-name()="anchor"]')
    assert len(anchors_neutral) == 0
    # Preserve the university name from metadata.
    assert meta_neutral.university in doc_neutral.paragraphs[1].text
    # Logo-free headers have no reserved left indent.
    p0_ind = doc_neutral.paragraphs[0]._p.xpath('.//w:ind')
    assert len(p0_ind) == 0 or p0_ind[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}left') in (None, "0")

    # Logo-enabled mode also takes the university name from metadata.
    meta_logo = ReportMetadata(
        university="ТЕСТОВЫЙ ТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ",
        logo_mode="default"
    )
    doc_logo = docx.Document(build_docx_report(meta_logo, "# Ход работы"))
    logo_text = "\n".join(p.text for p in doc_logo.paragraphs[:4])
    assert "ТЕСТОВЫЙ ТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ" in logo_text


def test_logo_mode_normalize_strict():
    """Verify normalize_logo_mode defaults unknown modes to none."""
    from app.engine.title_page import normalize_logo_mode
    assert normalize_logo_mode("default") == "default"
    assert normalize_logo_mode("custom") == "custom"
    assert normalize_logo_mode("none") == "none"
    assert normalize_logo_mode("samstu") == "none"
    assert normalize_logo_mode(None) == "none"
